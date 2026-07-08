#!/usr/bin/env python3
"""Render a .pptx to PNG previews with PIL (no LibreOffice needed).

Faithful enough to catch clutter, text overflow, and misalignment. Handles:
solid-fill rounded rects, borders, pictures (with alpha), text frames with
per-run font/size/color/bold/italic/mono + word wrap + v-anchor + alignment,
and tables. Not a pixel-perfect PowerPoint clone.
"""
import sys, io
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE_TYPE

SCALE = 200.0 / 914400.0          # EMU -> px  (200 px per inch)
PT2PX = 200.0 / 72.0              # pt -> px

FONTS = {
    (False, False, False): "/usr/share/fonts/liberation-sans/LiberationSans-Regular.ttf",
    (True,  False, False): "/usr/share/fonts/liberation-sans/LiberationSans-Bold.ttf",
    (False, True,  False): "/usr/share/fonts/liberation-sans/LiberationSans-Italic.ttf",
    (True,  True,  False): "/usr/share/fonts/liberation-sans/LiberationSans-BoldItalic.ttf",
    (False, False, True):  "/usr/share/fonts/dejavu/DejaVuSansMono.ttf",
    (True,  False, True):  "/usr/share/fonts/dejavu/DejaVuSansMono-Bold.ttf",
    (False, True,  True):  "/usr/share/fonts/dejavu/DejaVuSansMono-Oblique.ttf",
    (True,  True,  True):  "/usr/share/fonts/dejavu/DejaVuSansMono-BoldOblique.ttf",
}
_fc = {}
def font(pt, bold=False, italic=False, mono=False):
    px = max(6, int(round(pt * PT2PX)))
    key = (px, bold, italic, mono)
    if key not in _fc:
        _fc[key] = ImageFont.truetype(FONTS[(bool(bold), bool(italic), bool(mono))], px)
    return _fc[key]

def E(v):  # EMU -> px int
    return int(round((v or 0) * SCALE))

def rgb_of(color, default=None):
    try:
        if color is not None and color.type is not None:
            return "#" + str(color.rgb)
    except Exception:
        pass
    return default

def is_mono(name):
    if not name:
        return False
    n = name.lower()
    return "mono" in n or "consol" in n or "courier" in n or "menlo" in n or "dejavu sans mono" in n

def para_runs(p):
    """Yield (text, size_pt, bold, italic, mono, color) for each run."""
    out = []
    default_sz = 18
    for r in p.runs:
        sz = r.font.size.pt if r.font.size else (p.font.size.pt if p.font.size else default_sz)
        bold = bool(r.font.bold)
        italic = bool(r.font.italic)
        mono = is_mono(r.font.name)
        col = rgb_of(r.font.color, "#202124")
        out.append([r.text, sz, bold, italic, mono, col])
    return out

def wrap_runs(runs, maxw, line_spacing=1.12):
    """Greedy word-wrap a list of runs into lines. Returns list of lines;
    each line = list of (text, font, color, w, ascent, descent)."""
    # tokenize into words preserving trailing spaces per run
    tokens = []
    for text, sz, bold, italic, mono, col in runs:
        f = font(sz, bold, italic, mono)
        asc, desc = f.getmetrics()
        # split keeping spaces attached to the word before
        i = 0
        parts = text.replace("\n", " \n ").split(" ")
        for j, w in enumerate(parts):
            if w == "\n":
                tokens.append(("\n", f, col, 0, asc, desc))
                continue
            word = w + (" " if j < len(parts) - 1 else "")
            if word == "":
                continue
            tokens.append((word, f, col, f.getlength(word), asc, desc))
    lines, cur, curw = [], [], 0
    for tok in tokens:
        w = tok[0]
        if w == "\n":
            lines.append(cur); cur, curw = [], 0; continue
        tw = tok[3]
        if cur and curw + tw > maxw and w.strip():
            lines.append(cur); cur, curw = [], 0
        cur.append(tok); curw += tw
    if cur:
        lines.append(cur)
    return lines

def draw_text_frame(draw, tf, box, inset=(0.1, 0.05)):
    x0, y0, w, h = box
    il = int(inset[0] * 200); it = int(inset[1] * 200)
    # respect explicit margins if present
    try:
        if tf.margin_left is not None: il = E(tf.margin_left)
        if tf.margin_top is not None: it = E(tf.margin_top)
    except Exception:
        pass
    tx = x0 + il; tw = w - 2 * il
    # anchor
    anchor = tf.vertical_anchor
    # measure all paragraphs
    blocks = []
    total_h = 0
    for p in tf.paragraphs:
        runs = para_runs(p)
        if not runs:
            # empty paragraph -> blank line
            f = font(p.font.size.pt if p.font.size else 14)
            asc, desc = f.getmetrics()
            blocks.append(("blank", (asc + desc)))
            total_h += int((asc + desc) * 1.1)
            continue
        lines = wrap_runs(runs, tw)
        sb = int((p.space_before.pt if p.space_before else 0) * PT2PX)
        sa = int((p.space_after.pt if p.space_after else 0) * PT2PX)
        lh_list = []
        for ln in lines:
            asc = max((t[4] for t in ln), default=12)
            desc = max((t[5] for t in ln), default=4)
            lh = int((asc + desc) * 1.14)
            lh_list.append((asc, lh))
        ph = sb + sum(lh for _, lh in lh_list) + sa
        blocks.append(("para", p, lines, lh_list, sb, sa))
        total_h += ph
    if anchor == MSO_ANCHOR.MIDDLE:
        y = y0 + max(0, (h - total_h) // 2)
    elif anchor == MSO_ANCHOR.BOTTOM:
        y = y0 + max(0, h - total_h)
    else:
        y = y0 + it
    for b in blocks:
        if b[0] == "blank":
            y += int(b[1] * 1.1); continue
        _, p, lines, lh_list, sb, sa = b
        y += sb
        align = p.alignment
        for (ln, (asc, lh)) in zip(lines, lh_list):
            lw = sum(t[3] for t in ln)
            if align == PP_ALIGN.CENTER:
                lx = tx + max(0, (tw - lw) // 2)
            elif align == PP_ALIGN.RIGHT:
                lx = tx + max(0, tw - lw)
            else:
                lx = tx
            for (text, f, col, wpx, a, d) in ln:
                draw.text((lx, y + (asc - a)), text, font=f, fill=col or "#202124")
                lx += wpx
            y += lh
        y += sa

def rounded(draw, box, radius, fill=None, outline=None, width=2):
    x0, y0, w, h = box
    draw.rounded_rectangle([x0, y0, x0 + w, y0 + h], radius=radius,
                           fill=fill, outline=outline, width=width)

def render_shape(sh, img, draw):
    try:
        L, T, W, H = E(sh.left), E(sh.top), E(sh.width), E(sh.height)
    except Exception:
        L = T = W = H = 0
    box = (L, T, W, H)
    st = sh.shape_type
    # picture
    if st == MSO_SHAPE_TYPE.PICTURE:
        try:
            im = Image.open(io.BytesIO(sh.image.blob)).convert("RGBA")
            im = im.resize((max(1, W), max(1, H)))
            img.paste(im, (L, T), im)
        except Exception as e:
            draw.rectangle([L, T, L + W, T + H], outline="#cccccc")
        return
    # tables
    if sh.has_table:
        render_table(sh, draw); return
    # fill + border for autoshapes / textboxes
    fillc = None; linec = None; lw = 2
    try:
        if sh.fill.type == 1:  # solid
            fillc = "#" + str(sh.fill.fore_color.rgb)
    except Exception:
        pass
    try:
        if sh.line.color and sh.line.color.type is not None:
            linec = "#" + str(sh.line.color.rgb)
            if sh.line.width: lw = max(1, E(sh.line.width))
    except Exception:
        pass
    if fillc or linec:
        r = min(24, W // 2, H // 2) if (W > 40 and H > 40) else 0
        rounded(draw, box, r, fill=fillc, outline=linec, width=lw)
    if sh.has_text_frame and sh.text_frame.text.strip():
        draw_text_frame(draw, sh.text_frame, box)

def render_table(sh, draw):
    tbl = sh.table
    L, T = E(sh.left), E(sh.top)
    colw = [E(c.width) for c in tbl.columns]
    rowh = [E(r.height) for r in tbl.rows]
    y = T
    for ri, row in enumerate(tbl.rows):
        x = L
        for ci, cell in enumerate(row.cells):
            cw, ch = colw[ci], rowh[ri]
            fillc = None
            try:
                if cell.fill.type == 1:
                    fillc = "#" + str(cell.fill.fore_color.rgb)
            except Exception:
                pass
            draw.rectangle([x, y, x + cw, y + ch], fill=fillc, outline="#d0d0d0", width=1)
            if cell.text_frame.text.strip():
                draw_text_frame(draw, cell.text_frame, (x, y, cw, ch), inset=(0.08, 0.03))
            x += cw
        y += rowh[ri]

def render(pptx_path, out_dir, only=None):
    import os
    os.makedirs(out_dir, exist_ok=True)
    prs = Presentation(pptx_path)
    W = E(prs.slide_width); H = E(prs.slide_height)
    paths = []
    for i, slide in enumerate(prs.slides, 1):
        if only and i not in only:
            continue
        img = Image.new("RGB", (W, H), "white")
        draw = ImageDraw.Draw(img)
        for sh in slide.shapes:
            try:
                render_shape(sh, img, draw)
            except Exception as e:
                print(f"  slide{i}: shape error {e}", file=sys.stderr)
        p = os.path.join(out_dir, f"slide{i:02d}.png")
        img.save(p); paths.append(p)
    return paths, (W, H)

def contact_sheet(paths, out, cols=3, cellw=640, title_h=26):
    ims = [Image.open(p).convert("RGB") for p in paths]
    thumbs = []
    fnt = ImageFont.truetype(FONTS[(True, False, False)], 18)
    for p, im in zip(paths, ims):
        r = cellw / im.width
        th = im.resize((cellw, int(im.height * r)))
        canvas = Image.new("RGB", (cellw, th.height + title_h), "#e8e8e8")
        canvas.paste(th, (0, title_h))
        d = ImageDraw.Draw(canvas)
        d.rectangle([0, 0, cellw, title_h], fill="#0B2E6B")
        import os
        d.text((6, 4), os.path.basename(p), font=fnt, fill="white")
        thumbs.append(canvas)
    rows = (len(thumbs) + cols - 1) // cols
    rowh = [max(thumbs[r*cols+c].height for c in range(cols) if r*cols+c < len(thumbs)) for r in range(rows)]
    Wt = cellw * cols; Ht = sum(rowh)
    sheet = Image.new("RGB", (Wt, Ht), "#bbbbbb")
    y = 0
    for r in range(rows):
        x = 0
        for c in range(cols):
            k = r*cols+c
            if k < len(thumbs):
                sheet.paste(thumbs[k], (x, y))
            x += cellw
        y += rowh[r]
    sheet.save(out)
    return out

if __name__ == "__main__":
    pptx = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else "preview"
    only = None
    if len(sys.argv) > 3:
        only = set(int(x) for x in sys.argv[3].split(","))
    paths, size = render(pptx, out, only)
    print(f"rendered {len(paths)} slides at {size}")
    if len(paths) > 1:
        cs = contact_sheet(paths, out + "_sheet.png")
        print("contact sheet:", cs)
